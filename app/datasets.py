"""Planilhas que o usuário sobe (CSV ou Excel) viram uma lista de linhas {coluna: texto}."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from pathlib import Path

MAX_ROWS = 20000
MAX_BYTES = 30 * 1024 * 1024


class DatasetError(ValueError):
    pass


def _clean_header(cols: list) -> list[str]:
    out, seen = [], set()
    for i, c in enumerate(cols):
        name = re.sub(r"\s+", " ", str(c if c is not None else "")).strip() or f"coluna_{i + 1}"
        base, k = name, 2
        while name in seen:
            name = f"{base}_{k}"
            k += 1
        seen.add(name)
        out.append(name)
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v).strip()


def _with_ids(header: list[str], rows: list[list]) -> tuple[list[str], list[dict]]:
    """Garante uma coluna `id` única. Se a planilha não tiver, cria L0001, L0002..."""
    data = [dict(zip(header, [_cell(v) for v in r] + [""] * (len(header) - len(r)))) for r in rows]
    data = [d for d in data if any(d.values())]
    if not data:
        raise DatasetError("A planilha não tem linhas com dados.")
    if len(data) > MAX_ROWS:
        raise DatasetError(f"A planilha tem {len(data)} linhas. O limite da demo é {MAX_ROWS}.")
    ids = [d.get("id", "") for d in data]
    if "id" not in header or any(not i for i in ids) or len(set(ids)) != len(ids):
        if "id" in header:  # coluna id existe mas não serve: preserva o valor original
            header = ["id_original" if h == "id" else h for h in header]
            for d in data:
                d["id_original"] = d.pop("id")
        width = max(4, len(str(len(data))))
        for i, d in enumerate(data, 1):
            d["id"] = f"L{i:0{width}d}"
        header = ["id"] + header
        data = [{h: d.get(h, "") for h in header} for d in data]
    return header, data


def _read_csv(raw: bytes) -> tuple[list[str], list[list]]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DatasetError("Não consegui ler o CSV. Salve como UTF-8.")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = type("Fallback", (csv.excel,), {"delimiter": ";" if sample.count(";") > sample.count(",") else ","})
    rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise DatasetError("O CSV precisa de um cabeçalho e pelo menos uma linha.")
    return _clean_header(rows[0]), rows[1:]


def _read_xlsx(raw: bytes) -> tuple[list[str], list[list]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise DatasetError("Falta o pacote openpyxl. Rode: pip install -r requirements.txt") from e
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise DatasetError(f"Não consegui abrir o Excel: {e}") from e
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    wb.close()
    if len(rows) < 2:
        raise DatasetError("A primeira aba precisa de um cabeçalho e pelo menos uma linha.")
    width = max(len(r) for r in rows)
    header = rows[0] + [None] * (width - len(rows[0]))
    keep = [i for i in range(width) if header[i] not in (None, "") or any(i < len(r) and r[i] not in (None, "") for r in rows[1:])]
    return _clean_header([header[i] for i in keep]), [[r[i] if i < len(r) else None for i in keep] for r in rows[1:]]


def parse(filename: str, raw: bytes) -> tuple[list[str], list[dict]]:
    if len(raw) > MAX_BYTES:
        raise DatasetError("Arquivo maior que 30 MB.")
    ext = Path(filename or "").suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        header, rows = _read_xlsx(raw)
    elif ext in (".csv", ".txt", ".tsv"):
        header, rows = _read_csv(raw)
    elif ext == ".xls":
        raise DatasetError("Formato .xls antigo não é suportado. Salve como .xlsx ou .csv.")
    else:
        raise DatasetError("Envie um arquivo .csv ou .xlsx.")
    return _with_ids(header, rows)


def load_csv_file(path: Path) -> tuple[list[str], list[dict]]:
    return parse(path.name, path.read_bytes())
